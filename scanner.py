"""One scan cycle: for each bucket, fetch risk + de-escalation news, score it,
aggregate, and append one row per bucket to the running Excel log.

Run manually with: python scanner.py
Run on a schedule via .github/workflows/scan.yml (every 4 hours).
"""
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config import BUCKETS, MAX_ARTICLES_PER_QUERY
from news_fetcher import fetch_and_score_group
from sentiment_engine import score_text

LOG_PATH = Path(__file__).parent / "data" / "sentiment_log.xlsx"
LOG_SHEET = "Log"

HEADERS = [
    "timestamp_utc", "date", "bucket",
    "risk_articles", "risk_finbert_avg", "risk_vader_avg",
    "deescalation_articles", "deescalation_finbert_avg", "deescalation_vader_avg",
    "composite_score",
    "top_headline_1", "top_headline_1_url",
    "top_headline_2", "top_headline_2_url",
    "top_headline_3", "top_headline_3_url",
]


def _avg(values: list[float]) -> float:
    return sum(values) / len(values) if values else 0.0


def _composite_score(risk_count, risk_finbert_avg, deescalation_count, deescalation_finbert_avg) -> float:
    """
    Positive = net escalation pressure. Negative = net de-escalation. Zero = quiet on both sides.

    v1 heuristic, not a validated model -- reasonable starting weights to tune once
    you have real data to look at:
      risk_weight    = risk_count * (0.5 - risk_finbert_avg / 2)
                        (finbert_avg in [-1, 1]; more articles + more negative tone -> higher)
      deescal_weight = deescalation_count * (0.5 + deescalation_finbert_avg / 2)
                        (more articles + more positive tone -> higher, subtracted)

    This is the mechanism that lets "conflict cooled off" actually register as a
    negative (de-escalating) score instead of just flatlining at zero from silence.
    """
    risk_weight = risk_count * (0.5 - risk_finbert_avg / 2)
    deescalation_weight = deescalation_count * (0.5 + deescalation_finbert_avg / 2)
    return round(risk_weight - deescalation_weight, 3)


def _top_headlines(articles: list[dict], n: int = 3) -> list[dict]:
    ranked = sorted(articles, key=lambda a: abs(a["finbert"]), reverse=True)[:n]
    padded = ranked + [{"title": "", "link": ""}] * (n - len(ranked))
    return padded


def _ensure_workbook():
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    if LOG_PATH.exists():
        wb = load_workbook(LOG_PATH)
        if LOG_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(LOG_SHEET)
            ws.append(HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = LOG_SHEET
        ws.append(HEADERS)
    return wb


def run_scan():
    now = datetime.now(timezone.utc)
    timestamp = now.isoformat(timespec="seconds")
    date_str = now.strftime("%Y-%m-%d")

    wb = _ensure_workbook()
    ws = wb[LOG_SHEET]

    for bucket_key, bucket in BUCKETS.items():
        print(f"Scanning bucket: {bucket['label']}")

        risk_articles = fetch_and_score_group(bucket["risk_queries"], MAX_ARTICLES_PER_QUERY, score_text)
        deescalation_articles = fetch_and_score_group(
            bucket["deescalation_queries"], MAX_ARTICLES_PER_QUERY, score_text
        )

        risk_finbert_avg = _avg([a["finbert"] for a in risk_articles])
        risk_vader_avg = _avg([a["vader"] for a in risk_articles])
        deescalation_finbert_avg = _avg([a["finbert"] for a in deescalation_articles])
        deescalation_vader_avg = _avg([a["vader"] for a in deescalation_articles])

        composite = _composite_score(
            len(risk_articles), risk_finbert_avg,
            len(deescalation_articles), deescalation_finbert_avg,
        )

        top3 = _top_headlines(risk_articles, 3)

        row = [
            timestamp, date_str, bucket["label"],
            len(risk_articles), round(risk_finbert_avg, 3), round(risk_vader_avg, 3),
            len(deescalation_articles), round(deescalation_finbert_avg, 3), round(deescalation_vader_avg, 3),
            composite,
        ]
        for h in top3:
            row.extend([h.get("title", ""), h.get("link", "")])

        ws.append(row)
        print(
            f"  risk={len(risk_articles)} articles (avg finbert {risk_finbert_avg:.2f}), "
            f"deescalation={len(deescalation_articles)} articles, composite={composite}"
        )

    wb.save(LOG_PATH)
    print(f"Saved log -> {LOG_PATH}")


if __name__ == "__main__":
    run_scan()
